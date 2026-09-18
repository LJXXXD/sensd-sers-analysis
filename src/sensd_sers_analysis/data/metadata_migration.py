"""Plan and apply a reversible per-spectrum metadata migration.

Run with ``python -m sensd_sers_analysis.data.metadata_migration`` to inspect,
or append ``--apply`` after reviewing the report. Existing archival workbooks
are immutable; every source must match its archived copy before migration.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
from pathlib import Path

import openpyxl

from sensd_sers_analysis.config.metadata_migration import (
    BACKUP_DIRECTORY,
    CONFLICT_POLICY,
    PAA_INITIAL_TARGET,
    REPORT_PATH,
    SOURCE_DIRECTORY,
)
from sensd_sers_analysis.config.metadata_schema import INITIAL_TARGET_LABEL, SAMPLE_TYPE_LABEL

logger = logging.getLogger(__name__)
ROOT = Path(__file__).resolve().parents[3]


def classify_label(value: str) -> tuple[str, str] | None:
    """Extract explicit sample identity and treatment from legacy source text."""
    text = value.casefold()
    if "rinsate" in text:
        return "Rinsate control", ""
    if "untreated" in text:
        return "Bacteria sample", ""
    if "heat" in text:
        return "Bacteria sample", "Heat treated"
    match = re.search(r"(\d+)\s*ppm", text)
    if match:
        return "Bacteria sample", f"PAA {match.group(1)} ppm"
    if "cfu" in text:
        return "Bacteria sample", ""
    return None


def migrate_rows(rows: list[list], *, collection: str) -> tuple[list[list], list[dict]]:
    """Map legacy rows to explicit sample identities, retaining unknowns for review."""
    labels = {
        str(row[0]).strip().casefold(): i for i, row in enumerate(rows) if isinstance(row[0], str)
    }
    if "sample type" in labels and INITIAL_TARGET_LABEL.casefold() in labels:
        return rows, []
    target_idx = labels["target concentration (cfu/ml)"]
    actual_idx = labels["actual concentration (cfu/ml)"]
    treatment_idx = labels["special treatment"]
    file_idx = labels["file name"]
    types = [SAMPLE_TYPE_LABEL]
    treatments = ["Special Treatment"]
    targets = [INITIAL_TARGET_LABEL]
    changes = []
    for col in range(1, len(rows[file_idx])):
        label = str(rows[treatment_idx][col] or "")
        filename = str(rows[file_idx][col] or "")
        target = rows[target_idx][col]
        actual = rows[actual_idx][col]
        explicit = classify_label(label)
        named = classify_label(filename)
        conflict = explicit is not None and named is not None and explicit != named
        rationale = "Explicit legacy treatment" if explicit else "Source TXT filename"
        identity = explicit or named
        if conflict and CONFLICT_POLICY == "review":
            identity = ("", explicit[1])
            rationale = "REVIEW: treatment label conflicts with source TXT filename"
        elif conflict and CONFLICT_POLICY == "filename":
            identity = named
            rationale = "LJ-directed source TXT precedence"
        elif conflict:
            rationale = "LJ-directed embedded metadata precedence"
        if identity is None and collection not in {"PAA", "Heat kill"}:
            if target == 0 and actual == 0:
                identity = ("Rinsate control", "")
                rationale = "One-time legacy untreated control mapping (target and actual zero)"
            elif isinstance(target, (int, float)) and target > 0:
                identity = ("Bacteria sample", "")
                rationale = "One-time legacy untreated positive initial-target mapping"
        if identity is None:
            identity = ("", "")
            rationale = "REVIEW: insufficient identity evidence"
        sample_type, treatment = identity
        if sample_type == "Bacteria sample" and treatment.startswith("PAA") and target is None:
            target = PAA_INITIAL_TARGET
            rationale += "; PAA initial target confirmed by LJ"
        if sample_type == "Bacteria sample" and treatment == "Heat treated" and target is None:
            match = re.search(r"(\d+)\s*cfu", label, flags=re.IGNORECASE)
            if match:
                target = int(match.group(1))
                rationale += "; initial target from treatment text"
            else:
                rationale += "; REVIEW: Heat initial target not explicitly recorded"
        types.append(sample_type or None)
        treatments.append(treatment or None)
        targets.append(target)
        changes.append(
            {
                "column": col + 1,
                "source_txt_filename": filename,
                "old_treatment": rows[treatment_idx][col],
                "sample_type": sample_type,
                "special_treatment": treatment,
                "old_target": rows[target_idx][col],
                "initial_target": target,
                "actual": actual,
                "basis": rationale,
            }
        )
    output = []
    for idx, row in enumerate(rows):
        if idx == treatment_idx:
            output.extend([types, treatments])
        elif idx == target_idx:
            output.append(targets)
        else:
            output.append(row)
    return output, changes


def workbook_rows(path: Path) -> list[list]:
    """Read full worksheet values without changing their numerical precision."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"Expected one source sheet: {path}")
        return [list(row) for row in workbook.active.values]
    finally:
        workbook.close()


def spectral_rows(rows: list[list]) -> list[list]:
    """Extract the original Raman header and spectral matrix for exact comparison."""
    index = next(i for i, row in enumerate(rows) if str(row[0]).casefold() == "raman shift")
    return rows[index:]


def main() -> None:
    """Preflight every workbook, then optionally apply all prepared transformations."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    plans = []
    prepared = []
    for path in sorted((ROOT / SOURCE_DIRECTORY).rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        backup = ROOT / BACKUP_DIRECTORY / path.relative_to(ROOT / SOURCE_DIRECTORY)
        original = workbook_rows(path)
        migrated, changes = migrate_rows(original, collection=path.parent.name)
        if not changes:
            continue
        if path.read_bytes() != backup.read_bytes():
            raise ValueError(f"Source differs from archive; no files changed: {path}")
        if spectral_rows(original) != spectral_rows(migrated):
            raise ValueError(f"Unexpected spectral modification: {path}")
        plans.append(
            {
                "path": str(path.relative_to(ROOT)),
                "backup": str(backup.relative_to(ROOT)),
                "before_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "changes": changes,
            }
        )
        prepared.append((path, migrated))
    report = {"applied": False, "conflict_policy": CONFLICT_POLICY, "files": plans}
    report_path = ROOT / REPORT_PATH
    if not plans and report_path.exists():
        logger.info("All files already migrated; preserving original migration report.")
        return
    report_path.write_text(json.dumps(report, indent=2))
    if args.apply:
        staged = []
        try:
            for path, migrated in prepared:
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.title = "Sheet1"
                for row in migrated:
                    sheet.append(row)
                temporary = path.with_suffix(".migration.xlsx")
                staged.append((temporary, path))
                workbook.save(temporary)
                workbook.close()
                if workbook_rows(temporary) != migrated:
                    raise ValueError(f"Export verification failed: {path}")
            for temporary, path in staged:
                temporary.replace(path)
        finally:
            for temporary, _ in staged:
                temporary.unlink(missing_ok=True)
        report["applied"] = True
        for record, (path, _) in zip(plans, prepared, strict=True):
            record["after_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        report_path.write_text(json.dumps(report, indent=2))
    logger.info(
        "%s %d workbooks; report: %s",
        "Migrated" if args.apply else "Planned",
        len(plans),
        report_path,
    )


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    main()
