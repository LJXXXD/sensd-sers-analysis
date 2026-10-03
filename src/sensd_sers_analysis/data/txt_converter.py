"""Instrument TXT parsing, metadata values/presets, and embedded workbook serialization.

These functions operate on explicit values and bytes without Streamlit state.
"""

from __future__ import annotations

import re
from datetime import date, datetime, time
from io import BytesIO, StringIO
from typing import Any

import numpy as np
import pandas as pd

from sensd_sers_analysis.data.io import RAMAN_SHIFT_DECIMALS
from sensd_sers_analysis.config.metadata_schema import (
    INITIAL_TARGET_LABEL,
    SAMPLE_TYPE_LABEL,
    DATE_METADATA_FIELD_NUMBERS,
    TIME_METADATA_FIELD_NUMBERS,
    NUMERIC_METADATA_FIELD_NUMBERS,
    OPTIONAL_METADATA_FIELD_NUMBERS,
    METADATA_FIELD_SPECS,
    METADATA_LOGICAL_GROUPS,
    METADATA_WIDGET_KEYS,
)

TEMPLATE_VERSION = 1
TARGET_CONCENTRATION_LABEL = INITIAL_TARGET_LABEL
ACTUAL_CONCENTRATION_LABEL = "Actual Concentration (CFU/mL)"
FILE_NAME_LABEL = "File Name"
SPECIAL_TREATMENT_LABEL = "Special Treatment"
INTENSITY_HEADER = "Relative Light intensity (a.u)"
RAMAN_SHIFT_HEADER = "Raman Shift"
_SEPARATOR_AFTER_METADATA_NUMBERS = frozenset(group[-1] for group in METADATA_LOGICAL_GROUPS)


def parse_txt_content(content: str) -> pd.DataFrame:
    """
    Parse one instrument TXT export into Raman shift and intensity columns.

    Parameters
    ----------
    content:
        UTF-8 text of the instrument export. The initial ``Data from ... Node``
        provenance header, blank lines, and ``>`` comments are not measurements.

    Returns
    -------
    pd.DataFrame
        Columns ``RamanShift`` and ``Value``.
    """

    lines = content.splitlines()
    first = next((i for i, line in enumerate(lines) if line.strip()), None)
    if (
        first is not None
        and lines[first].strip().startswith("Data from ")
        and lines[first].strip().endswith(" Node")
    ):
        del lines[first]
    df = pd.read_csv(
        StringIO("\n".join(lines)),
        sep="\t",
        comment=">",
        header=None,
    )
    if df.shape[1] != 2 or df.empty:
        raise ValueError("TXT spectrum requires exactly two nonempty numeric columns.")
    df.columns = ["RamanShift", "Value"]
    df = df.apply(pd.to_numeric, errors="coerce")
    if not np.isfinite(df.to_numpy(dtype=float)).all():
        raise ValueError("TXT coordinates and intensities must be finite numeric measurements.")
    if df["RamanShift"].duplicated().any():
        raise ValueError("TXT Raman coordinates must be unique.")
    return df.reset_index(drop=True)


def validate_common_shift(
    file_contents: dict[str, str],
    *,
    atol: float = 1e-4,
) -> tuple[np.ndarray | None, str | None]:
    """
    Verify all uploaded TXT files share the same Raman shift grid.

    Returns
    -------
    tuple
        ``(common_shift, error_message)`` — shift array when valid, else ``None`` and message.
    """

    common_shift: np.ndarray | None = None
    for name, content in file_contents.items():
        try:
            shifts = parse_txt_content(content)["RamanShift"].to_numpy()
        except (ValueError, pd.errors.ParserError) as exc:
            return None, f"Invalid TXT file {name}: {exc}"
        if common_shift is None:
            common_shift = shifts
            continue
        if common_shift.shape != shifts.shape or not np.allclose(
            common_shift, shifts, atol=atol, rtol=0
        ):
            return None, f"Raman shift mismatch in file: {name}"
    return common_shift, None


def merge_txt_spectra(
    file_contents: dict[str, str],
    *,
    min_shift: float | None,
    max_shift: float | None,
) -> tuple[np.ndarray | None, list[np.ndarray] | None, str | None]:
    """
    Merge parsed TXT spectra into aligned Raman shift and intensity columns.

    Parameters
    ----------
    file_contents:
        Mapping of filename to UTF-8 file text, in column order.
    min_shift:
        Optional lower Raman shift bound (cm⁻¹).
    max_shift:
        Optional upper Raman shift bound (cm⁻¹).

    Returns
    -------
    tuple
        ``(raman_shift, intensity_columns, error_message)``.
    """

    names = list(file_contents.keys())
    if not names:
        return None, None, "No TXT files provided."

    common_shift, shift_error = validate_common_shift(file_contents)
    if shift_error:
        return None, None, shift_error
    if common_shift is None:
        return None, None, "No Raman shift data found."

    min_final = float(min_shift) if min_shift is not None else float(common_shift.min())
    max_final = float(max_shift) if max_shift is not None else float(common_shift.max())
    if not np.isfinite([min_final, max_final]).all():
        return None, None, "Raman shift bounds must be finite."
    if min_final >= max_final:
        return None, None, "Min Raman Shift must be less than Max."

    mask = (common_shift >= min_final) & (common_shift <= max_final)
    raman_shift = common_shift[mask]
    if raman_shift.size < 2:
        return None, None, "Raman shift range must retain at least two measured coordinates."
    intensity_columns: list[np.ndarray] = []

    for name in names:
        df = parse_txt_content(file_contents[name])
        trimmed = df.loc[mask, "Value"].to_numpy(dtype=float)
        if len(trimmed) != len(raman_shift):
            return None, None, f"Truncated row count mismatch for file: {name}"
        intensity_columns.append(trimmed)

    return raman_shift, intensity_columns, None


def parse_required_number(raw: str) -> float | None:
    """Parse a required numeric field; returns ``None`` when empty or non-numeric."""

    stripped = raw.strip()
    if not stripped:
        return None
    try:
        value = float(stripped)
        return value if np.isfinite(value) else None
    except ValueError:
        return None


def parse_optional_number(raw: str) -> float | None:
    """
    Parse an optional numeric field.

    Returns ``None`` when empty (allowed). Returns ``None`` when non-numeric
    (caller should treat as validation failure).
    """

    return parse_required_number(raw)


def to_excel_number(value: float) -> int | float:
    """Coerce a numeric value for Excel cells (integers stored without decimals)."""

    if value.is_integer():
        return int(value)
    return value


def parse_metadata_date(raw: str) -> date | None:
    """Parse a metadata date string in ISO or common US format."""

    stripped = raw.strip()
    if not stripped:
        return None
    try:
        return date.fromisoformat(stripped)
    except ValueError:
        pass
    for fmt in ("%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(stripped, fmt).date()
        except ValueError:
            continue
    return None


def normalize_ampm_period(raw_period: str) -> str:
    """Return ``AM`` or ``PM`` from a case-insensitive period token."""

    token = raw_period.upper().replace(".", "")
    return "AM" if token.startswith("A") else "PM"


def ampm_parse_candidates(raw: str) -> list[str]:
    """
    Build normalized 12-hour time strings with an explicit AM/PM suffix.

    Accepts compact input such as ``230pm`` or ``2:30 PM`` and expands it into
    ``strptime``-compatible candidates.
    """

    text = re.sub(r"\s+", " ", raw.strip())
    if not text:
        return []

    ampm_match = re.search(r"(a\.?m\.?|p\.?m\.?)\.?\s*$", text, re.IGNORECASE)
    if not ampm_match:
        return [text]

    period = normalize_ampm_period(ampm_match.group(1))
    time_part = text[: ampm_match.start()].strip()
    if not time_part:
        return []

    candidates: list[str] = []
    if ":" in time_part:
        candidates.append(f"{time_part} {period}")
    else:
        if not time_part.isdigit() or len(time_part) > 4:
            return []
        digits = time_part
        if len(digits) <= 2:
            candidates.append(f"{int(digits)} {period}")
        elif len(digits) == 3:
            candidates.append(f"{digits[0]}:{digits[1:]} {period}")
        else:
            candidates.append(f"{int(digits[:2])}:{digits[2:]} {period}")

    deduped: list[str] = []
    for candidate in candidates:
        if candidate not in deduped:
            deduped.append(candidate)
    return deduped


def parse_metadata_time(raw: str) -> time | None:
    """Parse a metadata time string in 12-hour AM/PM or legacy 24-hour format."""

    stripped = raw.strip()
    if not stripped:
        return None

    ampm_formats = ("%I:%M:%S %p", "%I:%M %p", "%I %p")
    for candidate in ampm_parse_candidates(stripped):
        for fmt in ampm_formats:
            if fmt == "%I %p" and ":" in candidate:
                continue
            if fmt.startswith("%I:%M") and candidate.count(":") != fmt.count(":"):
                continue
            try:
                return datetime.strptime(candidate, fmt).time()
            except ValueError:
                continue

    legacy_formats = ("%H:%M:%S", "%H:%M")
    for candidate in (stripped,):
        for fmt in legacy_formats:
            if fmt == "%H:%M" and len(candidate) != 5:
                continue
            try:
                return datetime.strptime(candidate, fmt).time()
            except ValueError:
                continue
    return None


def format_metadata_time_display(parsed: time) -> str:
    """Format a ``time`` value as ``HH:MM AM/PM`` with zero-padded hour and minute."""

    hour = parsed.hour % 12 or 12
    ampm = "AM" if parsed.hour < 12 else "PM"
    if parsed.second or parsed.microsecond:
        return f"{hour:02d}:{parsed.minute:02d}:{parsed.second:02d} {ampm}"
    return f"{hour:02d}:{parsed.minute:02d} {ampm}"


def format_metadata_date(value: date | str | None) -> str:
    """Return a normalized ``YYYY-MM-DD`` string, or empty when invalid."""

    if value is None:
        return ""
    if isinstance(value, date):
        return value.isoformat()
    parsed = parse_metadata_date(str(value))
    return parsed.isoformat() if parsed else ""


def format_metadata_time(value: time | str | None) -> str:
    """Return a normalized 12-hour AM/PM time string, or empty when invalid."""

    if value is None:
        return ""
    if isinstance(value, time):
        parsed = value
    else:
        parsed = parse_metadata_time(str(value))
    if parsed is None:
        return ""
    return format_metadata_time_display(parsed)


def coerce_metadata_widget_state_value(field_number: int, raw_value: Any) -> Any:
    """Convert stored template or snapshot text to a widget-ready session value."""

    if field_number in DATE_METADATA_FIELD_NUMBERS:
        if isinstance(raw_value, date):
            return raw_value
        if raw_value is None:
            return None
        stripped = str(raw_value).strip()
        return parse_metadata_date(stripped) if stripped else None
    if field_number in TIME_METADATA_FIELD_NUMBERS:
        if isinstance(raw_value, time):
            return raw_value
        if raw_value is None:
            return None
        stripped = str(raw_value).strip()
        return parse_metadata_time(stripped) if stripped else None
    if raw_value is None:
        return ""
    return str(raw_value).strip()


def metadata_widget_state_to_string(field_number: int, raw_value: Any) -> str:
    """Serialize a metadata widget session value to workbook/template text."""

    if field_number in DATE_METADATA_FIELD_NUMBERS:
        return format_metadata_date(raw_value)
    if field_number in TIME_METADATA_FIELD_NUMBERS:
        return format_metadata_time(raw_value)
    if raw_value is None:
        return ""
    return str(raw_value).strip()


def is_metadata_widget_value_empty(field_number: int, raw_value: Any) -> bool:
    """Return whether a metadata widget has no usable value."""

    return not metadata_widget_state_to_string(field_number, raw_value)


def field_number_for_widget_key(widget_key: str) -> int | None:
    """Return the metadata field number for a widget key, if known."""

    return next(
        (field_number for field_number, _, key in METADATA_FIELD_SPECS if key == widget_key),
        None,
    )


def is_template_field_exportable(field_number: int, raw_value: str) -> bool:
    """
    Return whether a metadata field has a valid, exportable value.

    Geometry/acquisition fields must be finite numbers. Date (13) and Testing Time (14)
    require valid dates/times. Optional Notes (16) may be omitted
    when empty. All other fields must be non-empty text.
    """

    stripped = raw_value.strip() if isinstance(raw_value, str) else raw_value
    if field_number in OPTIONAL_METADATA_FIELD_NUMBERS:
        if is_metadata_widget_value_empty(field_number, stripped):
            return False
        return True
    if field_number in NUMERIC_METADATA_FIELD_NUMBERS:
        return parse_required_number(str(stripped).strip()) is not None
    if field_number in DATE_METADATA_FIELD_NUMBERS:
        return format_metadata_date(stripped) != ""
    if field_number in TIME_METADATA_FIELD_NUMBERS:
        return format_metadata_time(stripped) != ""
    return bool(str(stripped).strip())


def serialize_template_field_value(
    field_number: int, raw_value: str | Any
) -> str | int | float | None:
    """Convert a raw widget value to a JSON-safe template value, or ``None`` when invalid."""

    if field_number in NUMERIC_METADATA_FIELD_NUMBERS:
        parsed = parse_required_number(str(raw_value).strip())
        if parsed is None:
            return None
        return to_excel_number(parsed)
    if field_number in DATE_METADATA_FIELD_NUMBERS:
        formatted = format_metadata_date(raw_value)
        return formatted if formatted else None
    if field_number in TIME_METADATA_FIELD_NUMBERS:
        formatted = format_metadata_time(raw_value)
        return formatted if formatted else None
    stripped = str(raw_value).strip()
    if not stripped:
        return None
    return stripped


def build_template_export_payload(
    field_values: dict[str, str],
    selected_keys: frozenset[str],
    signal_labels: dict[str, dict[str, str]] | None = None,
) -> dict[str, Any] | None:
    """
    Build a setup-template JSON payload from selected metadata widget keys.

    Only checked fields with valid values are included.
    """

    export_fields: dict[str, str | int | float] = {}
    for field_number, _, widget_key in METADATA_FIELD_SPECS:
        if widget_key not in selected_keys:
            continue
        serialized = serialize_template_field_value(field_number, field_values.get(widget_key, ""))
        if serialized is None:
            continue
        export_fields[widget_key] = serialized
    if not export_fields and not signal_labels:
        return None
    payload: dict[str, Any] = {"version": TEMPLATE_VERSION, "fields": export_fields}
    if signal_labels:
        payload["signal_labels"] = signal_labels
    return payload


def apply_template_import_to_values(
    payload: dict[str, Any],
    field_values: dict[str, str],
) -> tuple[dict[str, str], list[str]]:
    """
    Merge template payload values into a widget-key mapping.

    Returns
    -------
    tuple
        Updated values and human-readable warning messages.
    """

    warnings: list[str] = []
    if payload.get("version") != TEMPLATE_VERSION:
        warnings.append(f"Unsupported template version: {payload.get('version')!r}")
        return field_values, warnings

    raw_fields = payload.get("fields")
    if not isinstance(raw_fields, dict):
        warnings.append("Template is missing a valid 'fields' object.")
        return field_values, warnings

    updated = dict(field_values)
    for widget_key, raw_value in raw_fields.items():
        if widget_key not in METADATA_WIDGET_KEYS:
            warnings.append(f"Unknown template field ignored: {widget_key}")
            continue
        field_number = field_number_for_widget_key(widget_key)
        if field_number is None:
            continue
        if isinstance(raw_value, bool):
            warnings.append(f"Invalid value for {widget_key}; skipped.")
            continue
        if isinstance(raw_value, (int, float)):
            if field_number in DATE_METADATA_FIELD_NUMBERS | TIME_METADATA_FIELD_NUMBERS:
                warnings.append(f"Invalid date/time value for {widget_key}; skipped.")
                continue
            if field_number in NUMERIC_METADATA_FIELD_NUMBERS:
                parsed = parse_required_number(str(raw_value))
                if parsed is None:
                    warnings.append(f"Invalid numeric value for {widget_key}; skipped.")
                    continue
                updated[widget_key] = str(to_excel_number(parsed))
            else:
                updated[widget_key] = str(raw_value)
            continue
        if isinstance(raw_value, str):
            if field_number in NUMERIC_METADATA_FIELD_NUMBERS:
                parsed = parse_required_number(raw_value)
                if parsed is None:
                    warnings.append(f"Invalid numeric value for {widget_key}; skipped.")
                    continue
                updated[widget_key] = str(to_excel_number(parsed))
            elif field_number in DATE_METADATA_FIELD_NUMBERS:
                formatted = format_metadata_date(raw_value)
                if not formatted:
                    warnings.append(f"Invalid date value for {widget_key}; skipped.")
                    continue
                updated[widget_key] = formatted
            elif field_number in TIME_METADATA_FIELD_NUMBERS:
                formatted = format_metadata_time(raw_value)
                if not formatted:
                    warnings.append(f"Invalid time value for {widget_key}; skipped.")
                    continue
                updated[widget_key] = formatted
            else:
                updated[widget_key] = raw_value.strip()
            continue
        warnings.append(f"Unsupported value type for {widget_key}; skipped.")

    return updated, warnings


def build_embedded_workbook_rows(
    metadata: dict[str, Any],
    *,
    target_concentrations: list[int | float | str],
    actual_concentrations: list[int | float],
    source_txt_filenames: list[str],
    special_treatments: list[str],
    raman_shift: np.ndarray,
    intensity_columns: list[np.ndarray],
    sample_types: list[str] | None = None,
) -> list[list[Any]]:
    """
    Build row-major cell data for the embedded-metadata Excel layout.

    Parameters
    ----------
    metadata:
        Mapping of Excel column-A labels to column-B values (rows 1–16).
    target_concentrations:
        Target concentration labels per signal column.
    actual_concentrations:
        Measured concentration values per signal column.
    source_txt_filenames:
        Original instrument ``.txt`` filename per signal column (provenance).
    special_treatments:
        Optional treatment label per signal (e.g. heat kill, PAA); blank when unused.
    raman_shift:
        Raman shift grid (cm⁻¹).
    intensity_columns:
        One intensity vector per signal column.
    sample_types:
        Explicit sample identity per signal; blank identities require completion.

    Returns
    -------
    list[list[Any]]
        Rectangular rows suitable for ``DataFrame.to_excel(header=False)``.
    """

    n_signals = len(intensity_columns)
    if n_signals == 0:
        raise ValueError("Workbook export requires at least one spectrum.")
    aligned = [
        target_concentrations,
        actual_concentrations,
        source_txt_filenames,
        special_treatments,
    ]
    if sample_types is not None:
        aligned.append(sample_types)
    if any(len(values) != n_signals for values in aligned):
        raise ValueError("Per-signal metadata must match the number of spectra.")
    if any(
        np.asarray(signal).shape != np.asarray(raman_shift).shape for signal in intensity_columns
    ):
        raise ValueError("Each intensity vector must align with the Raman coordinates.")
    if not np.isfinite(raman_shift).all() or not all(
        np.isfinite(signal).all() for signal in intensity_columns
    ):
        raise ValueError("Workbook spectra require finite coordinates and intensities.")
    grid = np.asarray(raman_shift)
    if (
        grid.ndim != 1
        or grid.size < 1
        or np.unique(np.round(grid, RAMAN_SHIFT_DECIMALS)).size != grid.size
    ):
        raise ValueError(
            "Workbook Raman coordinates must be distinct at loader precision and contain at least one point."
        )
    actual = np.asarray(actual_concentrations, dtype=float)
    if not np.isfinite(actual).all() or np.any(actual < 0):
        raise ValueError("Actual concentrations must be finite nonnegative measured values.")
    for target in target_concentrations:
        if str(target).strip() and (
            parse_required_number(str(target)) is None or float(target) < 0
        ):
            raise ValueError("Initial targets must be blank or finite nonnegative numbers.")
    width = 1 + n_signals

    def _pad_row(row: list[Any]) -> list[Any]:
        return row + [""] * (width - len(row))

    rows: list[list[Any]] = [
        _pad_row([excel_label, metadata[excel_label]]) for _, excel_label, _ in METADATA_FIELD_SPECS
    ]
    rows.append(_pad_row([FILE_NAME_LABEL, *source_txt_filenames]))
    rows.append(_pad_row([SAMPLE_TYPE_LABEL, *(sample_types or [""] * n_signals)]))
    rows.append(_pad_row([SPECIAL_TREATMENT_LABEL, *special_treatments]))
    rows.append(_pad_row([TARGET_CONCENTRATION_LABEL, *target_concentrations]))
    rows.append(_pad_row([ACTUAL_CONCENTRATION_LABEL, *actual_concentrations]))
    rows.append([RAMAN_SHIFT_HEADER, *[INTENSITY_HEADER] * n_signals])

    for shift_idx, shift_value in enumerate(raman_shift):
        rows.append(
            [float(shift_value)]
            + [float(intensity_columns[signal_idx][shift_idx]) for signal_idx in range(n_signals)]
        )

    return rows


def apply_full_width_row_separator(
    worksheet: Any,
    row_idx: int,
    max_col: int,
    bottom_side: Any,
) -> None:
    """
    Draw a horizontal separator beneath ``row_idx`` across columns 1..``max_col``.

    Only the bottom edge is styled; no left, right, or top borders are added.
    """
    from openpyxl.styles import Border

    for col_idx in range(1, max_col + 1):
        worksheet.cell(row=row_idx, column=col_idx).border = Border(bottom=bottom_side)


def embedded_workbook_to_excel_bytes(rows: list[list[Any]]) -> bytes:
    """
    Serialize embedded workbook rows to styled ``.xlsx`` bytes.

    Applies bold labels and full-width horizontal separators between metadata
    sections, before the spectral table, and after the actual concentration row.
    """

    from openpyxl import Workbook
    from openpyxl.styles import Font, Side

    n_metadata = len(METADATA_FIELD_SPECS)
    n_signals = len(rows[n_metadata]) - 1
    data_width = 1 + n_signals
    file_name_row_idx = n_metadata + 1
    sample_type_row_idx = (
        next(i for i, row in enumerate(rows, 1) if row[0] == SAMPLE_TYPE_LABEL)
        if any(row[0] == SAMPLE_TYPE_LABEL for row in rows)
        else None
    )
    special_treatment_row_idx = next(
        i for i, row in enumerate(rows, 1) if row[0] == SPECIAL_TREATMENT_LABEL
    )
    target_row_idx = next(
        i for i, row in enumerate(rows, 1) if row[0] == TARGET_CONCENTRATION_LABEL
    )
    actual_row_idx = next(
        i for i, row in enumerate(rows, 1) if row[0] == ACTUAL_CONCENTRATION_LABEL
    )
    header_row_idx = next(i for i, row in enumerate(rows, 1) if row[0] == RAMAN_SHIFT_HEADER)

    separator_row_indices = [
        field_number
        for field_number, _, _ in METADATA_FIELD_SPECS
        if field_number in _SEPARATOR_AFTER_METADATA_NUMBERS
    ]
    separator_row_indices.append(actual_row_idx)

    thin_side = Side(style="thin", color="000000")
    bold_font = Font(bold=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"

    for row_idx, row_values in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_values, start=1):
            if value == "":
                continue
            worksheet.cell(row=row_idx, column=col_idx, value=value)

        if row_idx <= n_metadata:
            worksheet.cell(row=row_idx, column=1).font = bold_font
            continue

        if row_idx in {
            target_row_idx,
            actual_row_idx,
            file_name_row_idx,
            special_treatment_row_idx,
            sample_type_row_idx,
        }:
            worksheet.cell(row=row_idx, column=1).font = bold_font
            continue

        if row_idx == header_row_idx:
            for col_idx in range(1, data_width + 1):
                worksheet.cell(row=row_idx, column=col_idx).font = bold_font

    for separator_row_idx in separator_row_indices:
        apply_full_width_row_separator(worksheet, separator_row_idx, data_width, thin_side)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
